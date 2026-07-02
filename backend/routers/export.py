"""
Export endpoints — generate XLSX and PDF files from the current search results.
"""

import io
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import and_
from sqlalchemy.orm import Session

from database import get_db
from models import Voter
from routers.voters import _fts_search_ids

router = APIRouter()

# Shared column definitions for export
EXPORT_COLUMNS = [
    ("voter_id", "EPIC / Voter ID"),
    ("full_name", "Full Name"),
    ("father_or_husband_name", "Father / Husband Name"),
    ("age", "Age"),
    ("gender", "Gender"),
    ("house_door_no", "House / Door No"),
    ("state", "State"),
    ("district", "District"),
    ("parliamentary_constituency", "Parliamentary Constituency"),
    ("assembly_constituency", "Assembly Constituency"),
    ("part_no", "Part No"),
    ("part_name", "Part Name"),
    ("serial_no", "Serial No"),
    ("polling_station_no", "Polling Station No"),
    ("polling_station_name", "Polling Station Name"),
]


def _build_query(db: Session, q, voter_id, full_name, state, district,
                  parliamentary_constituency, assembly_constituency, part_no, gender):
    """Build a filtered SQLAlchemy query (shared between XLSX and PDF exports)."""
    query = db.query(Voter)

    if q and q.strip():
        ids = _fts_search_ids(q.strip(), limit=50_000)
        if not ids:
            return query.filter(False)
        query = query.filter(Voter.id.in_(ids))

    filters = []
    if voter_id:
        filters.append(Voter.voter_id.ilike(f"%{voter_id}%"))
    if full_name:
        filters.append(Voter.full_name.ilike(f"%{full_name}%"))
    if state:
        filters.append(Voter.state.ilike(f"%{state}%"))
    if district:
        filters.append(Voter.district.ilike(f"%{district}%"))
    if parliamentary_constituency:
        filters.append(Voter.parliamentary_constituency.ilike(f"%{parliamentary_constituency}%"))
    if assembly_constituency:
        filters.append(Voter.assembly_constituency.ilike(f"%{assembly_constituency}%"))
    if part_no is not None:
        filters.append(Voter.part_no == part_no)
    if gender:
        filters.append(Voter.gender.ilike(f"{gender}%"))

    if filters:
        query = query.filter(and_(*filters))

    return query.order_by(Voter.full_name)


@router.get("/xlsx")
async def export_xlsx(
    q: Optional[str] = Query(None),
    voter_id: Optional[str] = Query(None),
    full_name: Optional[str] = Query(None),
    state: Optional[str] = Query(None),
    district: Optional[str] = Query(None),
    parliamentary_constituency: Optional[str] = Query(None),
    assembly_constituency: Optional[str] = Query(None),
    part_no: Optional[int] = Query(None),
    gender: Optional[str] = Query(None),
    limit: int = Query(10_000, le=100_000),
    db: Session = Depends(get_db),
):
    """Export filtered voter records as an Excel (.xlsx) file."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    voters = _build_query(
        db, q, voter_id, full_name, state, district,
        parliamentary_constituency, assembly_constituency, part_no, gender
    ).limit(limit).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Voter Records"

    # ── Header styling ──
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [label for _, label in EXPORT_COLUMNS]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    ws.row_dimensions[1].height = 20

    # ── Data rows ──
    alt_fill = PatternFill("solid", fgColor="F0F4F8")
    for row_idx, voter in enumerate(voters, 2):
        fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, (field, _) in enumerate(EXPORT_COLUMNS, 1):
            val = getattr(voter, field, None)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border
            if fill:
                cell.fill = fill

    # ── Column widths ──
    col_widths = [15, 28, 28, 6, 8, 14, 18, 20, 30, 30, 8, 22, 8, 14, 30]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Freeze header row
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="voter_records.xlsx"'},
    )


@router.get("/pdf")
async def export_pdf(
    q: Optional[str] = Query(None),
    voter_id: Optional[str] = Query(None),
    full_name: Optional[str] = Query(None),
    state: Optional[str] = Query(None),
    district: Optional[str] = Query(None),
    parliamentary_constituency: Optional[str] = Query(None),
    assembly_constituency: Optional[str] = Query(None),
    part_no: Optional[int] = Query(None),
    gender: Optional[str] = Query(None),
    limit: int = Query(1_000, le=10_000),
    db: Session = Depends(get_db),
):
    """Export filtered voter records as a formatted PDF table."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    voters = _build_query(
        db, q, voter_id, full_name, state, district,
        parliamentary_constituency, assembly_constituency, part_no, gender
    ).limit(limit).all()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "Title", parent=styles["Title"], fontSize=14, textColor=colors.HexColor("#1E3A5F")
    )
    small = ParagraphStyle("Small", parent=styles["Normal"], fontSize=7)

    # Compact column set for PDF (landscape A4 has limited width)
    PDF_COLS = [
        ("serial_no", "#", 8 * mm),
        ("voter_id", "EPIC No", 22 * mm),
        ("full_name", "Name", 35 * mm),
        ("father_or_husband_name", "Father/Husband", 30 * mm),
        ("age", "Age", 8 * mm),
        ("gender", "Sex", 8 * mm),
        ("house_door_no", "House No", 18 * mm),
        ("assembly_constituency", "Assembly", 25 * mm),
        ("part_no", "Part", 8 * mm),
        ("polling_station_name", "Polling Station", 35 * mm),
    ]

    header = [col[1] for col in PDF_COLS]
    col_widths = [col[2] for col in PDF_COLS]

    data = [header]
    for v in voters:
        row = []
        for field, _, _ in PDF_COLS:
            val = getattr(v, field, "") or ""
            row.append(Paragraph(str(val), small) if len(str(val)) > 15 else str(val))
        data.append(row)

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTSIZE", (0, 1), (-1, -1), 7),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F8")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))

    content = [
        Paragraph("Indian Voter Data — Exported Records", title_style),
        Spacer(1, 4 * mm),
        Paragraph(f"Total records: {len(voters)}", styles["Normal"]),
        Spacer(1, 4 * mm),
        table,
    ]
    doc.build(content)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="voter_records.pdf"'},
    )
